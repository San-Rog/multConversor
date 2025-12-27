import streamlit as st
import pandas as pd
import os
import io
import csv
import zipfile
from io import BytesIO
from io import StringIO
import openpyxl
import odf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches

class messages():
    def __init__(self, *args):
        self.fileTmp = args[0]
        self.suffix = args[1]
        self.nFiles = args[2]
        if self.nFiles == 1:
            self.fileFinal = f'zipado_sozinho_{self.suffix}.zip'
        else:
            self.fileFinal = f'zipado_múltiplos_{self.suffix}.zip'
        if None not in args:
            self.mensResult()
    
    def mensResult(self):
        if self.nFiles <= 1:
            exprFile = ['arquivo não repetido', 'abri-lo']
        else:
            exprFile = ['arquivos não repetidos', 'abri-los']
        if self.suffix in ['tsv']:
            mensStr = f':blue[**{self.fileFinal}**] com ***{self.nFiles} {exprFile[0]}***. Download aqui 👉.\n' \
                      f'(Use **Bloco de Notas** ou aplicativo similar para {exprFile[1]}.)'
        else:
            mensStr = f':blue[**{self.fileFinal}**] com ***{self.nFiles} {exprFile[0]}***. Download aqui 👉.'
        colMens, colZip = st.columns([19, 3], width='stretch', vertical_alignment='center')
        colMens.success(mensStr, icon='✔️')                              
        with open(self.fileTmp, "rb") as file:
            buttDown = colZip.download_button(label='',
                                              data=file,
                                              file_name=self.fileFinal,
                                              mime='application/zip', 
                                              icon=':material/download:', 
                                              use_container_width=True, 
                                              key='buttDown')
     
    @st.dialog(' ')
    def configTwo(self, str):
        st.markdown(str)  
      
class downFiles():
    def __init__(self, *args):
        self.files = args[0]
        self.index = args[1]
        self.engine = args[2]
        self.ext = args[3]
        self.opt = args[4]
        self.filesZip = []
        self.nFiles = 0
        if self.index in [0, 1, 2]:
            if self.index in [0, 1]:
                self.csvXlsx() 
            elif self.index == 2:
                self.csvTsv()
        elif self.index == 3:
            self.csvDocx()     
        if self.opt is not None:
            self.nameZip = f'arquivo_all_{self.ext}.zip'
            self.downZip()
            if os.path.getsize(self.nameZip) > 0:
                messages(self.nameZip, self.ext, self.nFiles)
            
    def csvXlsx(self):
        for f, file in enumerate(self.files):
            self.nameFile = file[0]
            self.dataFile = file[1]
            if self.opt == 2:
                self.fileOut = f'{self.nameFile}.xlsx'
            else:
                self.fileOut = f'{self.nameFile}.{self.ext}'
            self.sheetName = 'aba_única'
            wb = openpyxl.Workbook()
            sheet = wb.active
            sheet.title = 'dados'        
            for data in self.dataFile:
                try:
                    newData = [item.encode('ISO-8859-1').decode('utf8') for item in data]
                except: 
                    newData = [item for item in data]
                sheet.append(newData)
            wb.save(self.fileOut)
            self.bytesFiles(0)
    
    def csvTsv(self):
        for f, file in enumerate(self.files):
            self.nameFile = file[0]
            self.dataFile = file[1]
            self.fileOut = f'{self.nameFile}_new.csv'
            self.csvCsv(0)
            self.bytesFiles(1)
               
    def csvPd(self):
        file = self.files[0]
        self.nameFile = file[0]
        self.dataFile = file[1]
        self.fileOut = self.nameFile
        self.csvCsv(1)
        self.df = pd.read_csv(self.fileOut).fillna('')
        self.renameHead()
        st.dataframe(self.df)
    
    def csvCsv(self, mode):
        allLines = []
        for data in self.dataFile:
            try:
                newData = [item.encode('ISO-8859-1').decode('utf8') for item in data]
            except: 
                newData = [item for item in data]
            allLines.append(newData)
        with open(self.fileOut, 'w', newline='', encoding='utf-8') as recordCsv:
            writerCsv = csv.writer(recordCsv)
            writerCsv.writerows(allLines)
        if mode == 1:
            return self.fileOut
    
    def csvDocx(self):
        for f, file in enumerate(self.files):
            self.nameFile = file[0]
            self.dataFile = file[1]
            self.fileOut = f'{self.nameFile}_new.csv'
            self.csvCsv(0)
            doc = Document()
            with open(self.fileOut, 'r', encoding='utf-8') as f:
                csv_reader = csv.reader(f)
                headers = next(csv_reader)
                num_cols = len(headers)
                table = doc.add_table(rows=1, cols=num_cols)
                table.style = 'Table Grid'
                table.page_width = Inches(11.7) 
                table.page_height = Inches(8.5)                 
                hdr_cells = table.rows[0].cells
                for i, header in enumerate(headers):
                    hdr_cells[i].text = header
                for row in csv_reader:
                    row_cells = table.add_row().cells
                    for i, cell_data in enumerate(row):
                        row_cells[i].text = cell_data
                self.fileOut = f'{self.nameFile}_new.{self.ext}'
                doc.save(self.fileOut)
                self.bytesFiles(2)
            
    def bytesFiles(self, mode):
        if mode == 0:
            if self.opt != 2:
                self.df = pd.read_excel(self.fileOut).fillna('')
                self.renameHead()    
                output = BytesIO()
                with pd.ExcelWriter(output, engine=self.engine) as writer:
                    self.df.to_excel(writer, index=False, sheet_name=self.sheetName)
                excelBytes = output.getvalue()
                zips = (self.fileOut, excelBytes)
                self.filesZip.append(zips) 
                self.nFiles += 1
            else:
                self.df = pd.read_excel(self.fileOut).fillna('')
                self.renameHead()
                output = BytesIO()
                with pd.ExcelWriter(output, engine=self.engine) as writer:
                    self.df.to_excel(writer, index=False, sheet_name=self.sheetName)
                self.fileOut = f'{self.nameFile}.{self.ext}'              
                xlsxBytes = output.getvalue()
                bytesObj = io.BytesIO(xlsxBytes)
                self.df = pd.read_excel(bytesObj).fillna('')
                self.renameHead()
                htmlStr = self.df.to_html(index=False)
                htmlBytes = htmlStr.encode('utf-8')
                zips = (self.fileOut, htmlBytes)
                self.filesZip.append(zips) 
                self.nFiles += 1
        elif mode == 1:
            self.df = pd.read_csv(self.fileOut).fillna('')
            self.renameHead()
            output = BytesIO()
            self.df.to_csv(output, sep='\t', index=False)
            csvBytes = output.getvalue()
            self.fileOut = f'{self.nameFile}.{self.ext}'
            zips = (self.fileOut, csvBytes)
            self.filesZip.append(zips) 
            self.nFiles += 1
        elif mode == 2:
            newDocx = f'{self.nameFile}.{self.ext}'
            output = BytesIO()
            with open(self.fileOut, 'rb') as arquivo:
                docxRead = arquivo.read()
            zips = (newDocx, docxRead)
            self.filesZip.append(zips) 
            self.nFiles += 1
    
    def renameHead(self):
        head = {}
        for col in self.df.columns:
            if col.lower().find('unnamed') >= 0:
                head[col] = ''
            else:
                head[col] = col
        self.df.rename(columns=head, inplace=True)
        
    def downZip(self):
        with zipfile.ZipFile(self.nameZip, 'w', zipfile.ZIP_DEFLATED) as zipf:
            for file in self.filesZip:
                nameFile = file[0]
                dataFile = file[1]
                zipf.writestr(nameFile, dataFile)
        
class main():
    def __init__(self):          
        st.set_page_config(initial_sidebar_state="collapsed", layout="wide")
        self.typeExt = sorted(['CSV', 'XLS', 'XLSX'])
        nIni = len(self.typeExt)
        self.typeExt.insert(0, '')
        colType, colUpload = st.columns([12, 17], width='stretch')
        self.keyUp = 'zero'
        self.keyFile = 'typeFile'
        with colType:
            with st.container(border=4, key='contType', gap='small', height="stretch"):
                helpBox = 'https://pt.wikipedia.org/wiki/Comma-separated_values, https://en.wikipedia.org/wiki/XLS' 
                helpBox += 'https://pt.wikipedia.org/wiki/Microsoft_Excel'
                st.markdown('❇️ Seleção de tipo + arrastamento/escolha de arquivos', unsafe_allow_html=True, 
                            text_alignment='center')
                self.typeFile = st.selectbox(f'📂 Tipos de arquivo ({nIni})', self.typeExt,
                                             help=f'Selecionar a extensão desejada. Para reiniciar, escolher a linha em branco. \n{helpBox}') 
                if not self.typeFile: 
                    upDisabled = True
                    self.typeStr = ''
                else:
                    upDisabled = False
                    self.typeStr = f':red[**{self.typeFile}**]'
                    st.space(size="small")                
                self.upLoad = st.file_uploader(f'📙 Arraste/escolha dois ou mais arquivos {self.typeStr}.', 
                                               type=self.typeFile, accept_multiple_files=True, key=self.keyUp, 
                                               disabled=upDisabled, 
                                               help='É integrado de todos os arquivos selecionados, mesmo que se repitam.' \
                                                    'No momento do download, contudo, serão tratados como se não se repetissem.') 
        with colUpload:  
            try:
                self.files = list(set([file.name for file in self.upLoad]))
            except:
                self.files = [] 
            if not self.typeFile:
                self.configImageEmpty(4)
            if self.typeFile:
                with st.container(border=4, key='contUpload', gap='small', height='stretch', 
                                  vertical_alignment='center'):
                    nUploads = len(self.upLoad)
                    if not self.typeFile:
                        self.configImageEmpty(None)                                                
                    else:
                        if not self.typeFile == self.typeExt[1]: 
                            self.configImageEmpty(None)
                        elif self.typeFile == self.typeExt[1]:
                            self.exts = {'openpyxl': ['xls', 'xlsx', 'html'], 'odf': ['ods'], 'tsv': ['tsv'], 
                                         'doc': ['docx'], 'yaml': ['yaml'], 'json': ['json'], 'xhtml': ['xhtml'],
                                         'toml': ['toml'], 'txt': ['txt'], 'pdf': ['pdf']}                        
                            self.newTypes = []
                            self.segregateTypes()
                            self.disableds = ['disabled' + str(w) for w in range(len(self.newTypes))]
                            if nUploads == 0:
                                self.setSessionState(True)
                            else:
                                self.setSessionState(False)
                            typeLow = self.typeFile.lower()
                            self.strFunc = ['Converter um ou mais arquivos', 'Convertendo']
                            self.stripe = f':red[**{self.typeFile.lower()}**]'
                            with st.container(border=None, key='contOne', gap='small', height='stretch', 
                                              vertical_alignment='center'):
                                nFiles = len(self.files)
                                if nFiles <= 0:
                                    titleSel = f'Arquivo selecionado ({nFiles})'
                                else:
                                    titleSel = f'Arquivos selecionados ({nFiles})'
                                if nFiles > 0:
                                    opts = sorted(self.files)
                                    opts.insert(0, '')
                                else:
                                    opts = []
                                buttOne, buttTwo, buttThree, buttFour, buttFive, buttSix = ['' for i in range(6)]
                                buttSeven, buttEight, buttNine, buttTen, buttEleven, buttTwelve = ['' for i in range(6)]
                                self.allButtons = [buttOne, buttTwo, buttThree, buttFour, buttFive, buttSix, 
                                                   buttSeven, buttEight, buttNine, buttTen, buttEleven, buttTwelve]
                                colOne, colTwo, colThree = st.columns(spec=3, width='stretch')
                                colFour, colFive, colSix = st.columns(spec=3, width='stretch')
                                colSeven, colEight, colNine = st.columns(spec=3, width='stretch')
                                colTen, colEleven, colTwelve = st.columns(spec=3, width='stretch')
                                self.colsButts = {0: [0, colOne, ':material/sync_alt:'], 1: [1, colTwo, ':material/swap_horiz:'], 
                                                  2: [2, colThree, ':material/table_convert:'], 3: [3, colFour, ':material/transform:'], 
                                                  4: [4, colFive, ':material/convert_to_text:'], 5: [5, colSix, ':material/edit_arrow_up:'], 
                                                  6: [6, colSeven, ':material/edit_arrow_up:'], 7: [7, colEight, ':material/edit_arrow_up:'], 
                                                  8: [8, colNine, ':material/edit_arrow_up:'], 9: [9, colTen, ':material/edit_arrow_up:'], 
                                                  10: [10, colEleven, ':material/edit_arrow_up:'], 11: [11, colTwelve, ':material/edit_arrow_up:']}
                                buttOne = self.setButtons(self.colsButts[0])
                                buttTwo = self.setButtons(self.colsButts[1])
                                buttThree = self.setButtons(self.colsButts[2]) 
                                buttFour = self.setButtons(self.colsButts[3])
                                buttFive = self.setButtons(self.colsButts[4])
                                buttSix = self.setButtons(self.colsButts[5])                                
                                buttSeven = self.setButtons(self.colsButts[6]) 
                                buttEight = self.setButtons(self.colsButts[7])
                                buttNine = self.setButtons(self.colsButts[8])
                                buttTen = self.setButtons(self.colsButts[9])
                                buttEleven = self.setButtons(self.colsButts[10])
                                buttTwelve = self.setButtons(self.colsButts[11])                                
                if self.upLoad:
                    nNotRep = len(self.files)
                    nRep = nUploads - nNotRep                                
                    exprLoad = self.singPlural(nUploads, 'escolhido', 'escolhidos')
                    exprNotRep = self.singPlural(nNotRep, 'não repetido', 'não repetidos')
                    exprRep = self.singPlural(nRep, 'repetido', 'repetidos')
                    with st.container(border=2, key='contRepNo', gap='small', height='stretch', 
                                      vertical_alignment='center'):
                        colTotal, colNotRep, colRep = st.columns(spec=3, width='stretch', vertical_alignment='center')
                        with colTotal.popover(f'Informações {nUploads}', icon='ℹ️', use_container_width=True): 
                            self.elem = st.selectbox(exprLoad, self.files)
                            if self.elem:
                                self.formatSelect()  
                        with colNotRep.popover(f'{nNotRep} {exprNotRep}', icon='👍', use_container_width=True):
                            st.text('mmmmmmmmmmm', width="content")                                        
                        with colRep.popover(f'{nRep} {exprRep}', icon='❌', use_container_width=True):
                            st.text('mmmmmmmmmmm', width="content")  
                    if any(self.allButtons):
                        ind = self.allButtons.index(True)
                        expr = f'{self.strFunc[1]} {nUploads} do formato {self.stripe} para o foramto {self.newTypes[ind]}...'
                        if self.allButtons[0]:                     
                            self.index = 0
                            self.opt = 0                                                                                
                        elif self.allButtons[1]:
                            self.index = 0
                            self.opt = 1
                        elif self.allButtons[2]:
                            self.index = 0
                            self.opt = 2                        
                        elif self.allButtons[3]:
                            self.index = 1
                            self.opt = 0 
                        elif self.allButtons[4]: 
                            self.index = 2
                            self.opt = 0 
                        elif self.allButtons[5]: 
                            self.index = 3
                            self.opt = 0 
                        elif self.allButtons[6]:
                            pass
                        elif self.allButtons[7]:
                            pass
                        elif self.allButtons[8]:
                            pass
                        elif self.allButtons[9]:
                            pass
                        elif self.allButtons[10]:
                            pass
                        elif self.allButtons[11]:
                            pass
                        try:
                            with st.spinner(expr):
                                self.keys = list(self.exts.keys())
                                self.key = self.keys[self.index]
                                self.values = self.exts[self.key]
                                self.ext = self.values[self.opt] 
                                self.filesRead = [] 
                                self.segregateFiles()
                                downFiles(self.filesRead, self.index, self.key, self.ext, self.opt)
                        except:
                            pass
    
    def formatSelect(self):
        st.markdown(f"""
        <style>
            .st-e4 {{max-width: {700}px !important;}} 
        </style>
        """, unsafe_allow_html=True)
    
    def singPlural(self, *args):
        if args[0] <= 1: 
            expr = args[1]
        else:
            expr = args[2]
        return expr
    
    def setButtons(self, elems):
        #[0, colOne, ':material/sync_alt:']
        n = elems[0]
        col = elems[1]
        ico = elems[2]
        butt = f'butt{n}'
        self.allButtons[n] = col.button(label=f'{self.stripe} para {self.newTypes[n]}', key=butt,
                                          use_container_width=True, 
                                          icon=ico, 
                                          disabled=st.session_state[self.disableds[n]],
                                          help=f'{self.strFunc[0]} {self.stripe} para {self.newTypes[n]}.') 
    
    def segregateTypes(self):
        listTypes = list(self.exts.values())
        for tipo in listTypes:
            self.newTypes += tipo
        self.newTypes = [f':red[**{new}**]' for new in self.newTypes]
    
    def configImageEmpty(self, border):
        with st.container(border=border, key='contZero', gap='small', height='stretch'):
            st.markdown(f'0️⃣  seleção de tipo e/ou arquivo', text_alignment='center') 
            st.image(''zero.jpg', use_container_width='stretch') 
    
    def setSessionState(self, state):
        for disabled in self.disableds:
            if disabled not in st.session_state:
                st.session_state[disabled] = True 
            else:
                st.session_state[disabled] = state
  
    def segregateFiles(self):
        filesFind = {}
        for upLoad in self.upLoad: 
            nameGlobal = upLoad.name
            nameFile, ext = os.path.splitext(nameGlobal)
            filesFind.setdefault(nameGlobal, 0)
            if nameGlobal in self.files:
                filesFind[nameGlobal] += 1
            if filesFind[nameGlobal] > 1:
                continue
            newName = f'{nameFile}.{self.ext}'
            dataBytes = upLoad.getvalue()
            dataString = dataBytes.decode('ISO-8859-1')
            self.fileMemory = io.StringIO(dataString)
            sep = self.detectSep()
            readerCsv = csv.reader(self.fileMemory, delimiter=sep)
            joinNameRead = (nameFile, readerCsv)
            self.filesRead.append(joinNameRead)
            
    def detectSep(self):
        lines = 1024*10
        sample = self.fileMemory.read(lines)
        self.fileMemory.seek(0)
        dialect = csv.Sniffer().sniff(sample)
        return dialect.delimiter
            
if __name__ == '__main__':
    #formatos adicionais: yaml, json, xhtml, toml
    with open('configCss.css') as f:
        css = f.read()
    st.markdown(f'<style>{css}</style>', unsafe_allow_html=True) 
    main()
